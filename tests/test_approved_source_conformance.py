import json
from pathlib import Path

import pytest

from ev_charging.bronze_ingestion import iter_kba_rows
from ev_charging.identity import sha256_file
from ev_charging.source_contracts import inspect_bnetza, inspect_destatis, inspect_kba


ROOT = Path(__file__).resolve().parents[1]


def test_exact_approved_source_revisions_conform():
    config = json.loads((ROOT / "config/source_revisions_v1.json").read_text(encoding="utf-8"))
    required = [
        ROOT / source["filename"]
        for source in config["sources"]
        if source["classification"] == "CORE_V1"
    ]
    if not all(path.is_file() for path in required):
        pytest.skip("Official source files are not redistributed in the public repository")
    inspectors = {"bnetza": inspect_bnetza, "kba": inspect_kba, "destatis": inspect_destatis}
    observed = {}
    for source in config["sources"]:
        if source["classification"] != "CORE_V1":
            continue
        path = ROOT / source["filename"]
        assert path.stat().st_size == source["size_bytes"]
        assert sha256_file(path) == source["sha256"]
        observed[source["provider"]] = inspectors[source["provider"]](
            path, source["expected_reference_date"]
        )

    assert observed["bnetza"]["facility_rows"] == 116443
    assert observed["bnetza"]["distinct_facility_ids"] == 116443
    assert observed["kba"]["analytical_rows"] == 400
    assert observed["kba"]["distinct_statistical_keys"] == 400
    assert observed["kba"]["integer_bev_rows"] == 400
    assert observed["destatis"]["accepted_hierarchy_rows"] == 15984
    assert observed["destatis"]["district_rows"] == 401
    assert observed["destatis"]["distinct_district_codes"] == 401


def test_bnetza_contract_fails_on_physical_schema_change(tmp_path):
    path = tmp_path / "changed.csv"
    path.write_text(
        "Letzte Aktualisierung vom: 01.09.2026\n" + "\n" * 9 + "unexpected_header\nvalue\n",
        encoding="utf-8-sig",
    )
    with pytest.raises(ValueError, match="expected 47 columns"):
        inspect_bnetza(path, "2026-09-01")


def test_kba_analytical_row_selection_uses_five_digit_key(tmp_path):
    from datetime import datetime

    from openpyxl import Workbook

    path = tmp_path / "selection.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FZ 27.15"
    sheet.cell(13, 2, "Hessen")
    sheet.cell(13, 3, "06415")
    sheet.cell(13, 4, "HANAU")
    sheet.cell(14, 2, "")
    sheet.cell(14, 3, "Insgesamt")
    sheet.cell(14, 4, "TOTAL")
    workbook.save(path)
    metadata = {
        "source_revision_id": "revision",
        "source_file": path.name,
        "source_sha256": "a" * 64,
        "reference_date": "2026-07-01",
        "ingestion_run_id": "run",
        "ingested_at": datetime(2026, 9, 13),
    }

    rows = list(iter_kba_rows(path, metadata))
    assert len(rows) == 1
    assert rows[0][1] == "06415"
