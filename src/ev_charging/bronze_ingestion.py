"""Source-controlled immutable landing to managed-Delta Bronze ingestion."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator

from openpyxl import load_workbook

from ev_charging.identity import canonical_json, date_set_id, sha256_file, source_revision_id
from ev_charging.source_contracts import (
    DESTATIS_COLUMNS,
    KBA_COLUMNS,
    inspect_bnetza,
    inspect_destatis,
    inspect_kba,
    normalize_excel_code,
)
from ev_charging.runtime import catalog_name, volume_root


CATALOG = catalog_name()
VOLUME_ROOT = volume_root()
TECHNICAL_COLUMNS = [
    "_source_revision_id",
    "_source_file",
    "_source_sha256",
    "_reference_date",
    "_source_row_number",
    "_ingestion_run_id",
    "_ingested_at",
]

KBA_DELTA_COLUMNS = [
    "land",
    "statistische_kennziffer",
    "zulassungsbezirk",
    "pkw_anzahl_insgesamt",
    "alternativer_antrieb_anzahl_insgesamt",
    "alternativer_antrieb_anteil_prozent",
    "elektro_antriebe_anzahl_insgesamt",
    "elektro_antriebe_anteil_prozent",
    "elektro_bev",
    "plug_in_hybrid",
    "hybrid_anzahl_insgesamt",
    "benzin_hybrid",
    "diesel_hybrid",
    "gas_insgesamt",
]

DESTATIS_DELTA_COLUMNS = [
    "satzart",
    "textkennzeichen",
    "ars_land",
    "ars_rb",
    "ars_kreis",
    "ars_vb",
    "ars_gem",
    "gemeindename",
    "flaeche_km2",
    "bevoelkerung_insgesamt",
    "bevoelkerung_maennlich",
    "bevoelkerung_weiblich",
    "bevoelkerung_je_km2",
    "postleitzahl_verwaltungssitz",
    "laengengrad",
    "breitengrad",
    "reisegebiet_schluessel",
    "reisegebiet_bezeichnung",
    "verstaedterung_schluessel",
    "verstaedterung_bezeichnung",
]

TABLES = {
    "bnetza": f"{CATALOG}.bronze.bnetza_ladesaeulenregister",
    "kba": f"{CATALOG}.bronze.kba_fz27_15",
    "destatis": f"{CATALOG}.bronze.destatis_gv_isys_hierarchy",
}


def raw_excel_text(value: Any) -> str:
    """Return a stable raw-cell representation without interpreting markers as zero."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def safe_column_name(header: str) -> str:
    text = header.casefold().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")
    text = "".join(character for character in unicodedata.normalize("NFKD", text) if not unicodedata.combining(character))
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    if not text or text[0].isdigit():
        text = f"source_{text}"
    return text


def bnetza_column_mapping(headers: list[str]) -> dict[str, str]:
    mapping = {header: safe_column_name(header) for header in headers}
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("BNetzA official headers collide after physical-name normalization")
    return mapping


def quarantine_balances(source_records: int, accepted_records: int, quarantined_records: int) -> bool:
    return source_records == accepted_records + quarantined_records


def idempotent_action(existing_revision_rows: int, expected_revision_rows: int) -> str:
    if existing_revision_rows == 0:
        return "INGEST"
    if existing_revision_rows == expected_revision_rows:
        return "NO_OP"
    raise ValueError(
        f"Existing immutable revision row count {existing_revision_rows} != expected {expected_revision_rows}"
    )


def landing_path(source: dict[str, Any]) -> str:
    return (
        f"{VOLUME_ROOT}/{source['provider']}/{source['expected_reference_date']}/"
        f"{source['sha256']}/{source['original_filename']}"
    )


def manifest_landing_path(identifier: str) -> str:
    return f"{VOLUME_ROOT}/manifests/{identifier}.json"


def iter_bnetza_rows(path: Path, mapping: dict[str, str], metadata: dict[str, Any]) -> Iterator[tuple[Any, ...]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for _ in range(10):
            handle.readline()
        reader = csv.DictReader(handle, delimiter=";", quotechar='"')
        headers = reader.fieldnames or []
        if list(mapping) != headers:
            raise ValueError("BNetzA physical header changed between validation and extraction")
        for physical_row, row in enumerate(reader, start=12):
            source_values = tuple(row[header] for header in headers)
            yield source_values + metadata_tuple(metadata, physical_row)


def iter_kba_rows(path: Path, metadata: dict[str, Any]) -> Iterator[tuple[Any, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["FZ 27.15"]
    for physical_row, values in enumerate(
        sheet.iter_rows(min_row=13, min_col=2, max_col=15, values_only=True), start=13
    ):
        try:
            key = normalize_excel_code(values[1], 5)
        except ValueError:
            continue
        if not re.fullmatch(r"\d{5}", key):
            continue
        raw_values = [raw_excel_text(value) for value in values]
        raw_values[1] = key
        yield tuple(raw_values) + metadata_tuple(metadata, physical_row)


def iter_destatis_rows(path: Path, metadata: dict[str, Any]) -> Iterator[tuple[Any, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Onlineprodukt_Gemeinden30062026"]
    accepted = {"10", "20", "30", "40", "50", "60"}
    for physical_row, values in enumerate(
        sheet.iter_rows(min_row=7, min_col=1, max_col=20, values_only=True), start=7
    ):
        try:
            satzart = normalize_excel_code(values[0], 2)
        except ValueError:
            continue
        if satzart not in accepted:
            continue
        raw_values = [raw_excel_text(value) for value in values]
        raw_values[0] = satzart
        yield tuple(raw_values) + metadata_tuple(metadata, physical_row)


def metadata_tuple(metadata: dict[str, Any], physical_row: int) -> tuple[Any, ...]:
    return (
        metadata["source_revision_id"],
        metadata["source_file"],
        metadata["source_sha256"],
        date.fromisoformat(metadata["reference_date"]),
        physical_row,
        metadata["ingestion_run_id"],
        metadata["ingested_at"],
    )


def _spark_schema(source_columns: list[str]) -> Any:
    from pyspark.sql.types import DateType, LongType, StringType, StructField, StructType, TimestampType

    fields = [StructField(name, StringType(), True) for name in source_columns]
    fields.extend(
        [
            StructField("_source_revision_id", StringType(), False),
            StructField("_source_file", StringType(), False),
            StructField("_source_sha256", StringType(), False),
            StructField("_reference_date", DateType(), False),
            StructField("_source_row_number", LongType(), False),
            StructField("_ingestion_run_id", StringType(), False),
            StructField("_ingested_at", TimestampType(), False),
        ]
    )
    return StructType(fields)


def _write_staging_batches(
    spark: Any,
    rows: Iterable[tuple[Any, ...]],
    source_columns: list[str],
    staging_table: str,
    batch_size: int = 5000,
) -> int:
    schema = _spark_schema(source_columns)
    batch: list[tuple[Any, ...]] = []
    written = 0
    first = True
    for row in rows:
        batch.append(row)
        if len(batch) < batch_size:
            continue
        mode = "overwrite" if first else "append"
        spark.createDataFrame(batch, schema).write.format("delta").mode(mode).saveAsTable(staging_table)
        written += len(batch)
        batch.clear()
        first = False
    if batch:
        mode = "overwrite" if first else "append"
        spark.createDataFrame(batch, schema).write.format("delta").mode(mode).saveAsTable(staging_table)
        written += len(batch)
        first = False
    if first:
        raise ValueError("Contract parser produced no Bronze rows")
    return written


def _append_table_row(spark: Any, table: str, row: dict[str, Any]) -> None:
    schema = spark.table(table).schema
    ordered = tuple(row.get(field.name) for field in schema.fields)
    spark.createDataFrame([ordered], schema).write.format("delta").mode("append").saveAsTable(table)


def ensure_audit_tables(spark: Any) -> None:
    statements = [
        """
        CREATE TABLE IF NOT EXISTS {catalog}.audit.source_revision (
          source_revision_id STRING, source_name STRING, provider STRING, original_filename STRING,
          reference_date DATE, retrieval_date DATE, sha256 STRING, size_bytes BIGINT,
          landing_path STRING, source_page STRING, source_url STRING, ingestion_eligibility BOOLEAN,
          registered_at TIMESTAMP, contract_version STRING, schema_fingerprint STRING, record_fingerprint STRING
        ) USING DELTA
        """,
        """
        CREATE TABLE IF NOT EXISTS {catalog}.audit.date_set_registry (
          date_set_id STRING, canonical_manifest_json STRING, kba_reference_date DATE,
          bnetza_reference_date DATE, destatis_reference_date DATE, created_at TIMESTAMP,
          state STRING, date_set_label STRING, record_fingerprint STRING
        ) USING DELTA
        """,
        """
        CREATE TABLE IF NOT EXISTS {catalog}.audit.ingestion_run (
          ingestion_run_id STRING, source_revision_id STRING, started_at TIMESTAMP,
          completed_at TIMESTAMP, status STRING, source_row_count BIGINT, accepted_row_count BIGINT,
          quarantined_row_count BIGINT, failure_reason STRING, bronze_destination STRING
        ) USING DELTA
        """,
        """
        CREATE TABLE IF NOT EXISTS {catalog}.audit.bronze_quarantine (
          ingestion_run_id STRING, source_revision_id STRING, source_row_number BIGINT,
          failure_rule STRING, raw_problematic_value STRING, quarantined_at TIMESTAMP
        ) USING DELTA
        """,
        """
        CREATE TABLE IF NOT EXISTS {catalog}.audit.bronze_reconciliation (
          ingestion_run_id STRING, source_revision_id STRING, source_row_count BIGINT,
          bronze_row_count BIGINT, quarantine_row_count BIGINT, distinct_natural_keys BIGINT,
          source_sha256 STRING, controls_json STRING, physical_schema_mapping_json STRING,
          reconciled_at TIMESTAMP, result STRING
        ) USING DELTA
        """,
    ]
    for statement in statements:
        spark.sql(statement.format(catalog=CATALOG))


def _register_source_revision(
    spark: Any, source: dict[str, Any], observed: dict[str, Any], revision: str, landed: str, now: datetime
) -> str:
    schema_payload = {
        "physical_schema": observed["physical_schema"],
        "reference_date": observed["reference_date"],
        "row_grain": observed["row_grain"],
    }
    schema_fingerprint = hashlib.sha256(canonical_json(schema_payload).encode("utf-8")).hexdigest()
    immutable_payload = {
        "source_revision_id": revision,
        "source_name": source["dataset_id"],
        "provider": source["provider_name"],
        "original_filename": source["original_filename"],
        "reference_date": source["expected_reference_date"],
        "retrieval_date": source["retrieval_date"],
        "sha256": source["sha256"],
        "size_bytes": source["size_bytes"],
        "landing_path": landed,
        "source_page": source["source_page"],
        "source_url": source["source_url"],
        "ingestion_eligibility": True,
        "contract_version": "phase1_source_contract_v1",
        "schema_fingerprint": schema_fingerprint,
    }
    fingerprint = hashlib.sha256(canonical_json(immutable_payload).encode("utf-8")).hexdigest()
    existing = spark.table(f"{CATALOG}.audit.source_revision").where(
        f"source_revision_id = {sql_literal(revision)}"
    ).collect()
    if existing:
        if len(existing) != 1 or existing[0]["record_fingerprint"] != fingerprint:
            raise ValueError(f"Immutable source-revision registry conflict for {revision}")
        return "NO_OP"
    row = {
        **immutable_payload,
        "reference_date": date.fromisoformat(source["expected_reference_date"]),
        "retrieval_date": date.fromisoformat(source["retrieval_date"]),
        "registered_at": now,
        "record_fingerprint": fingerprint,
    }
    _append_table_row(spark, f"{CATALOG}.audit.source_revision", row)
    return "REGISTERED"


def _register_date_set(spark: Any, manifest: dict[str, Any], now: datetime) -> tuple[str, str]:
    identifier, canonical = date_set_id(manifest)
    fingerprint = hashlib.sha256(canonical.encode("utf-8")).hexdigest()
    existing = spark.table(f"{CATALOG}.audit.date_set_registry").where(
        f"date_set_id = {sql_literal(identifier)}"
    ).collect()
    if existing:
        if len(existing) != 1 or existing[0]["canonical_manifest_json"] != canonical:
            raise ValueError(f"Immutable date-set registry conflict for {identifier}")
        return identifier, "NO_OP"
    _append_table_row(
        spark,
        f"{CATALOG}.audit.date_set_registry",
        {
            "date_set_id": identifier,
            "canonical_manifest_json": canonical,
            "kba_reference_date": date(2026, 7, 1),
            "bnetza_reference_date": date(2026, 9, 1),
            "destatis_reference_date": date(2026, 6, 30),
            "created_at": now,
            "state": "candidate",
            "date_set_label": "Core V1 source set — 2026 Q3",
            "record_fingerprint": fingerprint,
        },
    )
    return identifier, "REGISTERED"


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def _source_controls(source: str, observed: dict[str, Any]) -> tuple[int, int]:
    if source == "bnetza":
        return observed["facility_rows"], observed["distinct_facility_ids"]
    if source == "kba":
        return observed["analytical_rows"], observed["distinct_statistical_keys"]
    return observed["accepted_hierarchy_rows"], observed["distinct_district_codes"]


def _natural_key_count(spark: Any, source: str, table: str, revision: str) -> int:
    frame = spark.table(table).where(f"_source_revision_id = {sql_literal(revision)}")
    if source == "bnetza":
        return frame.select("ladeeinrichtungs_id").distinct().count()
    if source == "kba":
        return frame.select("statistische_kennziffer").distinct().count()
    return frame.where("satzart = '40'").selectExpr("concat(ars_land, ars_rb, ars_kreis) AS district_code").distinct().count()


def _source_row_iterator(
    source_name: str, path: Path, observed: dict[str, Any], metadata: dict[str, Any]
) -> tuple[list[str], dict[str, str], Iterable[tuple[Any, ...]]]:
    if source_name == "bnetza":
        mapping = bnetza_column_mapping(observed["physical_schema"])
        return list(mapping.values()), mapping, iter_bnetza_rows(path, mapping, metadata)
    if source_name == "kba":
        return KBA_DELTA_COLUMNS, dict(zip(KBA_COLUMNS, KBA_DELTA_COLUMNS)), iter_kba_rows(path, metadata)
    return (
        DESTATIS_DELTA_COLUMNS,
        dict(zip(DESTATIS_COLUMNS, DESTATIS_DELTA_COLUMNS)),
        iter_destatis_rows(path, metadata),
    )


def _validate_landed_source(source: dict[str, Any]) -> tuple[Path, str, dict[str, Any]]:
    path = Path(landing_path(source))
    if not path.is_file():
        raise FileNotFoundError(f"Landed source is missing: {path}")
    digest = sha256_file(path)
    if path.stat().st_size != source["size_bytes"] or digest != source["sha256"]:
        raise ValueError(f"Landed source bytes do not match approved revision: {path}")
    revision = source_revision_id(
        source["provider"], source["dataset_id"], source["expected_reference_date"], digest
    )
    inspectors = {"bnetza": inspect_bnetza, "kba": inspect_kba, "destatis": inspect_destatis}
    observed = inspectors[source["provider"]](path, source["expected_reference_date"])
    return path, revision, observed


def _publish_revision(
    spark: Any,
    source: dict[str, Any],
    path: Path,
    revision: str,
    observed: dict[str, Any],
    run_id: str,
    started: datetime,
) -> dict[str, Any]:
    from pyspark.sql import functions as F

    source_name = source["provider"]
    destination = TABLES[source_name]
    expected_rows, expected_natural_keys = _source_controls(source_name, observed)
    existing_rows = 0
    if spark.catalog.tableExists(destination):
        existing_rows = spark.table(destination).where(F.col("_source_revision_id") == revision).count()
    action = idempotent_action(existing_rows, expected_rows)
    metadata = {
        "source_revision_id": revision,
        "source_file": source["original_filename"],
        "source_sha256": source["sha256"],
        "reference_date": source["expected_reference_date"],
        "ingestion_run_id": run_id,
        "ingested_at": started,
    }
    source_columns, physical_mapping, rows = _source_row_iterator(source_name, path, observed, metadata)
    staging = f"{CATALOG}.bronze._phase2b_stage_{source_name}_{run_id.replace('-', '')}"
    try:
        if action == "INGEST":
            written = _write_staging_batches(spark, rows, source_columns, staging)
            if written != expected_rows or spark.table(staging).count() != expected_rows:
                raise ValueError(f"Staging count mismatch for {revision}")
            if not spark.catalog.tableExists(destination):
                spark.sql(f"CREATE TABLE {destination} USING DELTA AS SELECT * FROM {staging} WHERE 1 = 0")
            spark.sql(f"INSERT INTO {destination} SELECT * FROM {staging}")
        bronze_rows = spark.table(destination).where(F.col("_source_revision_id") == revision).count()
        natural_keys = _natural_key_count(spark, source_name, destination, revision)
        if bronze_rows != expected_rows or natural_keys != expected_natural_keys:
            raise ValueError(
                f"Bronze reconciliation failed for {revision}: rows={bronze_rows}/{expected_rows}, "
                f"keys={natural_keys}/{expected_natural_keys}"
            )
        quarantine_rows = 0
        if not quarantine_balances(expected_rows, bronze_rows, quarantine_rows):
            raise ValueError(f"Quarantine accounting failed for {revision}")
        completed = datetime.now(timezone.utc).replace(tzinfo=None)
        _append_table_row(
            spark,
            f"{CATALOG}.audit.ingestion_run",
            {
                "ingestion_run_id": run_id,
                "source_revision_id": revision,
                "started_at": started,
                "completed_at": completed,
                "status": "IDEMPOTENT_NO_OP" if action == "NO_OP" else "SUCCESS",
                "source_row_count": expected_rows,
                "accepted_row_count": bronze_rows,
                "quarantined_row_count": quarantine_rows,
                "failure_reason": None,
                "bronze_destination": destination,
            },
        )
        controls_json = canonical_json(observed)
        _append_table_row(
            spark,
            f"{CATALOG}.audit.bronze_reconciliation",
            {
                "ingestion_run_id": run_id,
                "source_revision_id": revision,
                "source_row_count": expected_rows,
                "bronze_row_count": bronze_rows,
                "quarantine_row_count": quarantine_rows,
                "distinct_natural_keys": natural_keys,
                "source_sha256": source["sha256"],
                "controls_json": controls_json,
                "physical_schema_mapping_json": canonical_json(physical_mapping),
                "reconciled_at": completed,
                "result": "PASS",
            },
        )
        return {
            "action": action,
            "bronze_row_count": bronze_rows,
            "destination": destination,
            "distinct_natural_keys": natural_keys,
            "quarantine_row_count": quarantine_rows,
            "source": source_name,
            "source_revision_id": revision,
        }
    finally:
        spark.sql(f"DROP TABLE IF EXISTS {staging}")


def _record_failed_run(
    spark: Any,
    run_id: str,
    revision: str,
    started: datetime,
    destination: str,
    error: Exception,
) -> None:
    _append_table_row(
        spark,
        f"{CATALOG}.audit.ingestion_run",
        {
            "ingestion_run_id": run_id,
            "source_revision_id": revision,
            "started_at": started,
            "completed_at": datetime.now(timezone.utc).replace(tzinfo=None),
            "status": "FAILED",
            "source_row_count": None,
            "accepted_row_count": 0,
            "quarantined_row_count": 0,
            "failure_reason": f"{type(error).__name__}: {error}"[:8000],
            "bronze_destination": destination,
        },
    )


def run_phase2b(spark: Any, config_path: str, manifest_path: str) -> dict[str, Any]:
    """Validate landed bytes/contracts, register identities and publish source-faithful Bronze."""
    config = json.loads(Path(config_path).read_text(encoding="utf-8"))
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    ensure_audit_tables(spark)
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    identifier, date_set_registry_action = _register_date_set(spark, manifest, now)
    expected_manifest_path = Path(manifest_landing_path(identifier))
    if not expected_manifest_path.is_file():
        raise FileNotFoundError(f"Canonical manifest is not landed at {expected_manifest_path}")
    if expected_manifest_path.read_bytes() != Path(manifest_path).read_bytes():
        raise ValueError("Landed canonical manifest file differs from source-controlled manifest")

    results = []
    for source in config["sources"]:
        if source["classification"] != "CORE_V1":
            continue
        started = datetime.now(timezone.utc).replace(tzinfo=None)
        run_id = str(uuid.uuid4())
        revision = source_revision_id(
            source["provider"], source["dataset_id"], source["expected_reference_date"], source["sha256"]
        )
        try:
            path, observed_revision, observed = _validate_landed_source(source)
            if observed_revision != revision:
                raise ValueError(f"Observed source revision {observed_revision} != configured {revision}")
            registry_action = _register_source_revision(
                spark, source, observed, revision, landing_path(source), started
            )
            result = _publish_revision(spark, source, path, revision, observed, run_id, started)
            result["source_registry_action"] = registry_action
            result["controls"] = observed
            results.append(result)
        except Exception as error:
            _record_failed_run(spark, run_id, revision, started, TABLES[source["provider"]], error)
            raise
    return {
        "date_set_id": identifier,
        "date_set_registry_action": date_set_registry_action,
        "manifest_landing_path": str(expected_manifest_path),
        "result": "PASS",
        "sources": results,
    }
