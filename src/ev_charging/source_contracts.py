"""Strict parsers for the three approved Core V1 physical sources."""

from __future__ import annotations

import csv
import re
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BNETZA_REQUIRED_COLUMNS = [
    "Ladeeinrichtungs-ID",
    "Betreiber",
    "Status",
    "Art der Ladeeinrichtung",
    "Anzahl Ladepunkte",
    "Nennleistung Ladeeinrichtung [kW]",
    "Inbetriebnahmedatum",
    "Bundesland",
    "Kreis/kreisfreie Stadt",
    "Straße",
    "Hausnummer",
    "Postleitzahl",
    "Ort",
    "Breitengrad",
    "Längengrad",
]
for _slot in range(1, 7):
    BNETZA_REQUIRED_COLUMNS.extend(
        [
            f"Steckertypen{_slot}",
            f"Nennleistung Stecker{_slot}",
            f"EVSE-ID{_slot}",
            f"Public Key{_slot}",
        ]
    )

DESTATIS_COLUMNS = [
    "Satzart",
    "Textkennzeichen",
    "ARS Land",
    "ARS RB",
    "ARS Kreis",
    "ARS VB",
    "ARS Gem",
    "Gemeindename",
    "Fläche km2",
    "Bevölkerung insgesamt",
    "Bevölkerung männlich",
    "Bevölkerung weiblich",
    "Bevölkerung je km2",
    "Postleitzahl Verwaltungssitz",
    "Längengrad",
    "Breitengrad",
    "Reisegebiet Schlüssel",
    "Reisegebiet Bezeichnung",
    "Verstädterung Schlüssel",
    "Verstädterung Bezeichnung",
]

KBA_COLUMNS = [
    "Land",
    "Statistische Kennziffer",
    "Zulassungsbezirk",
    "Anzahl insgesamt",
    "Alternativer Antrieb — Anzahl insgesamt",
    "Alternativer Antrieb — Anteil in %",
    "Elektro-Antriebe — Anzahl insgesamt",
    "Elektro-Antriebe — Anteil in %",
    "Elektro (BEV)",
    "Plug-in-Hybrid",
    "Hybrid — Anzahl insgesamt",
    "Benzin-Hybrid",
    "Diesel-Hybrid",
    "Gas insgesamt",
]

_GERMAN_NUMBER = re.compile(r"^(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d+)?$")
_FIVE_DIGIT_KEY = re.compile(r"^\d{5}$")
_BNETZA_DATE = re.compile(r"Letzte Aktualisierung vom:\s*(\d{2})\.(\d{2})\.(\d{4})")
_DESTATIS_DATE = re.compile(r"\bam\s+(\d{2})\.(\d{2})\.(\d{4})\b")
_KBA_DATE = re.compile(r"\bam\s+(\d{1,2})\.\s+(Januar|April|Juli|Oktober)\s+(\d{4})\b")
_KBA_MONTHS = {"Januar": 1, "April": 4, "Juli": 7, "Oktober": 10}


def normalize_excel_code(value: Any, width: int) -> str:
    if value is None or str(value).strip() == "":
        return ""
    if isinstance(value, bool):
        raise ValueError(f"Boolean cannot be an administrative code: {value!r}")
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2]
    if not text.isdigit() or len(text) > width:
        raise ValueError(f"Invalid width-{width} code: {value!r}")
    return text.zfill(width)


def destatis_district_code(land: Any, rb: Any, kreis: Any) -> str:
    return normalize_excel_code(land, 2) + normalize_excel_code(rb, 1) + normalize_excel_code(kreis, 2)


def parse_german_decimal(value: Any) -> Decimal:
    text = "" if value is None else str(value).strip()
    if not _GERMAN_NUMBER.fullmatch(text):
        raise ValueError(f"Invalid German decimal: {value!r}")
    try:
        return Decimal(text.replace(".", "").replace(",", "."))
    except InvalidOperation as error:
        raise ValueError(f"Invalid German decimal: {value!r}") from error


def parse_nonnegative_integer(value: Any) -> int:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Expected integer, observed {value!r}")
    try:
        parsed = Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation as error:
        raise ValueError(f"Expected integer, observed {value!r}") from error
    if parsed < 0 or parsed != parsed.to_integral_value():
        raise ValueError(f"Expected non-negative integer, observed {value!r}")
    return int(parsed)


def point_slot_populated(connector_value: Any) -> bool:
    return connector_value is not None and bool(str(connector_value).strip())


def point_max_power(value: Any) -> Decimal:
    tokens = [token.strip() for token in str(value).split(";") if token.strip()]
    if not tokens:
        raise ValueError(f"Populated point has no power token: {value!r}")
    return max(parse_german_decimal(token) for token in tokens)


def inspect_bnetza(path: Path, expected_reference_date: str) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        preamble = [handle.readline().rstrip("\r\n") for _ in range(10)]
        matches = _BNETZA_DATE.findall("\n".join(preamble))
        if len(matches) != 1:
            raise ValueError(f"Expected one BNetzA snapshot date, observed {matches!r}")
        day, month, year = matches[0]
        reference_date = date(int(year), int(month), int(day)).isoformat()
        if reference_date != expected_reference_date:
            raise ValueError(f"BNetzA reference date {reference_date} != {expected_reference_date}")

        reader = csv.DictReader(handle, delimiter=";", quotechar='"')
        headers = reader.fieldnames or []
        if len(headers) != 47:
            raise ValueError(f"BNetzA expected 47 columns, observed {len(headers)}")
        missing = [name for name in BNETZA_REQUIRED_COLUMNS if name not in headers]
        if missing:
            raise ValueError(f"BNetzA missing required columns: {missing}")

        facility_ids: list[str] = []
        statuses: Counter[str] = Counter()
        facility_rows = 0
        declared_points = 0
        populated_points = 0
        normal_points = 0
        fast_points = 0
        point_mismatches = 0
        all_status_power = Decimal("0")
        maintenance_power = Decimal("0")
        in_service_power = Decimal("0")

        for physical_row, row in enumerate(reader, start=12):
            if None in row:
                raise ValueError(f"BNetzA shifted row at physical row {physical_row}")
            facility_rows += 1
            facility_id = row["Ladeeinrichtungs-ID"].strip()
            if not facility_id:
                raise ValueError(f"BNetzA blank facility ID at physical row {physical_row}")
            facility_ids.append(facility_id)
            status = row["Status"].strip()
            statuses[status] += 1
            declared = parse_nonnegative_integer(row["Anzahl Ladepunkte"])
            if declared < 1 or declared > 6:
                raise ValueError(f"BNetzA declared point count outside 1..6 at row {physical_row}")
            populated = 0
            for slot in range(1, 7):
                if not point_slot_populated(row[f"Steckertypen{slot}"]):
                    continue
                populated += 1
                maximum = point_max_power(row[f"Nennleistung Stecker{slot}"])
                if maximum <= 22:
                    normal_points += 1
                else:
                    fast_points += 1
            declared_points += declared
            populated_points += populated
            if declared != populated:
                point_mismatches += 1

            power = parse_german_decimal(row["Nennleistung Ladeeinrichtung [kW]"])
            if power <= 0:
                raise ValueError(f"BNetzA non-positive facility power at row {physical_row}")
            all_status_power += power
            if status == "In Wartung":
                maintenance_power += power
            elif status == "In Betrieb":
                in_service_power += power

    if len(set(facility_ids)) != facility_rows:
        raise ValueError("BNetzA facility identifier is not unique within the revision")
    if point_mismatches:
        raise ValueError(f"BNetzA declared/populated point mismatch in {point_mismatches} rows")
    if statuses != Counter({"In Betrieb": 116423, "In Wartung": 20}):
        raise ValueError(f"BNetzA status distribution changed: {dict(statuses)}")
    expected = {
        "facility_rows": 116443,
        "declared_points": 209136,
        "populated_points": 209136,
        "normal_points": 154740,
        "fast_points": 54396,
        "all_status_power_kw": Decimal("9086797.5"),
        "maintenance_power_kw": Decimal("484.0"),
        "in_service_power_kw": Decimal("9086313.5"),
    }
    observed = {
        "facility_rows": facility_rows,
        "declared_points": declared_points,
        "populated_points": populated_points,
        "normal_points": normal_points,
        "fast_points": fast_points,
        "all_status_power_kw": all_status_power,
        "maintenance_power_kw": maintenance_power,
        "in_service_power_kw": in_service_power,
    }
    if observed != expected:
        raise ValueError(f"BNetzA control mismatch: observed={observed}, expected={expected}")
    return {
        "source": "bnetza",
        "reference_date": reference_date,
        "physical_schema": headers,
        "physical_column_count": len(headers),
        "preamble_physical_rows": [1, 10],
        "header_physical_row": 11,
        "data_first_physical_row": 12,
        "row_grain": "one Ladeeinrichtung per source revision",
        "natural_key": ["source_revision_id", "Ladeeinrichtungs-ID"],
        "required_columns_present": True,
        "facility_rows": facility_rows,
        "distinct_facility_ids": len(set(facility_ids)),
        "status_distribution": dict(sorted(statuses.items())),
        "declared_points": declared_points,
        "populated_points": populated_points,
        "normal_points": normal_points,
        "fast_points": fast_points,
        "point_count_mismatches": point_mismatches,
        "all_status_nominal_power_kw": str(all_status_power),
        "maintenance_nominal_power_kw": str(maintenance_power),
        "in_service_nominal_power_kw": str(in_service_power),
        "parsing_ambiguity_changes_grain": False,
    }


def inspect_kba(path: Path, expected_reference_date: str) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "FZ 27.15" not in workbook.sheetnames:
        raise ValueError("KBA workbook is missing exact sheet FZ 27.15")
    sheet = workbook["FZ 27.15"]
    title = sheet["B6"].value
    matches = _KBA_DATE.findall(str(title))
    if len(matches) != 1:
        raise ValueError(f"Expected one KBA title date, observed {matches!r}")
    day, month_name, year = matches[0]
    reference_date = date(int(year), _KBA_MONTHS[month_name], int(day)).isoformat()
    if reference_date != expected_reference_date:
        raise ValueError(f"KBA reference date {reference_date} != {expected_reference_date}")

    required_headers = {
        "C8": "Statistische \nKennziffer",
        "D8": "Zulassungsbezirk",
        "E8": "Anzahl insgesamt",
        "H10": "Elektro-Antriebe (ohne Brennstoffzelle (Wasserstoff))",
        "H11": "Anzahl insgesamt",
        "I11": "Anteil in %",
        "J12": "Elektro (BEV)",
        "K12": "Plug-in-Hybrid",
    }
    for cell, expected in required_headers.items():
        observed = sheet[cell].value
        if observed != expected:
            raise ValueError(f"KBA header {cell} changed: {observed!r} != {expected!r}")

    header_cells = {
        f"{chr(64 + column)}{row}": sheet.cell(row, column).value
        for row in range(8, 13)
        for column in range(2, 16)
        if sheet.cell(row, column).value is not None
    }
    keys: list[str] = []
    bev_markers = Counter()
    integer_bev_rows = 0
    analytical_rows = 0
    source_row_numbers: list[int] = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=13, min_col=2, max_col=15, values_only=True), start=13
    ):
        raw_key = row[1]
        try:
            key = normalize_excel_code(raw_key, 5)
        except ValueError:
            continue
        if not _FIVE_DIGIT_KEY.fullmatch(key):
            continue
        analytical_rows += 1
        source_row_numbers.append(row_number)
        keys.append(key)
        bev = row[8]
        if bev is None or str(bev).strip() in {"", ".", "-", "/"}:
            bev_markers["<blank>" if bev is None or str(bev).strip() == "" else str(bev).strip()] += 1
        else:
            parse_nonnegative_integer(bev)
            integer_bev_rows += 1
        for offset in (3, 9):
            parse_nonnegative_integer(row[offset])

    if analytical_rows != 400:
        raise ValueError(f"KBA expected 400 analytical rows, observed {analytical_rows}")
    if len(set(keys)) != 400:
        raise ValueError(f"KBA expected 400 unique keys, observed {len(set(keys))}")
    if integer_bev_rows != 400 or bev_markers:
        raise ValueError(f"KBA BEV contract changed: integer={integer_bev_rows}, markers={dict(bev_markers)}")
    return {
        "source": "kba",
        "reference_date": reference_date,
        "sheet": "FZ 27.15",
        "title_cell": "B6",
        "title": title,
        "physical_header_range": "B8:O12",
        "data_first_physical_row": 13,
        "physical_schema": KBA_COLUMNS,
        "observed_nonblank_header_cells": header_cells,
        "row_grain": "one Zulassungsbezirk per source revision and reference date",
        "natural_key": ["source_revision_id", "Statistische Kennziffer"],
        "required_columns_present": True,
        "analytical_rows": analytical_rows,
        "distinct_statistical_keys": len(set(keys)),
        "five_digit_keys": sum(bool(_FIVE_DIGIT_KEY.fullmatch(key)) for key in keys),
        "integer_bev_rows": integer_bev_rows,
        "bev_marker_or_blank_counts": dict(bev_markers),
        "first_analytical_physical_row": min(source_row_numbers),
        "last_analytical_physical_row": max(source_row_numbers),
        "parsing_ambiguity_changes_grain": False,
    }


def inspect_destatis(path: Path, expected_reference_date: str) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = "Onlineprodukt_Gemeinden30062026"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Destatis workbook is missing exact sheet {sheet_name}")
    sheet = workbook[sheet_name]
    title = sheet["A1"].value
    matches = _DESTATIS_DATE.findall(str(title))
    if len(matches) != 1:
        raise ValueError(f"Expected one Destatis title date, observed {matches!r}")
    day, month, year = matches[0]
    reference_date = date(int(year), int(month), int(day)).isoformat()
    if reference_date != expected_reference_date:
        raise ValueError(f"Destatis reference date {reference_date} != {expected_reference_date}")

    required_headers = {
        "A3": "Satzart",
        "B3": "Textkennzeichen",
        "C3": "Amtlicher Regionalschlüssel (ARS)",
        "C4": "Land",
        "D4": "RB",
        "E4": "Kreis",
        "F4": "VB",
        "G4": "Gem",
        "H3": "Gemeindename",
    }
    for cell, expected in required_headers.items():
        observed = sheet[cell].value
        if observed != expected:
            raise ValueError(f"Destatis header {cell} changed: {observed!r} != {expected!r}")
    if any(sheet.cell(6, column).value is not None for column in range(1, 21)):
        raise ValueError("Destatis expected physical row 6 to be blank")

    accepted = {"10", "20", "30", "40", "50", "60"}
    widths = {3: 2, 4: 1, 5: 2, 6: 4, 7: 3}
    accepted_rows = 0
    satzart_counts: Counter[str] = Counter()
    district_codes: list[str] = []
    trailing_metadata: list[dict[str, Any]] = []
    metadata_rows_started = False
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=7, min_col=1, max_col=20, values_only=True), start=7
    ):
        raw_satzart = row[0]
        if raw_satzart is None or str(raw_satzart).strip() == "":
            continue
        try:
            satzart = normalize_excel_code(raw_satzart, 2)
        except ValueError:
            metadata_rows_started = True
            trailing_metadata.append({"physical_row": row_number, "column_a": str(raw_satzart)})
            continue
        if satzart not in accepted:
            raise ValueError(f"Destatis unexpected code-like Satzart {satzart!r} at physical row {row_number}")
        if metadata_rows_started:
            raise ValueError(f"Destatis accepted Satzart resumes after footer metadata at physical row {row_number}")
        accepted_rows += 1
        satzart_counts[satzart] += 1
        codes = {column: normalize_excel_code(row[column - 1], width) for column, width in widths.items()}
        if not row[7]:
            raise ValueError(f"Destatis blank Gemeindename at physical row {row_number}")
        if satzart == "40":
            if not all(codes[column] for column in (3, 4, 5)):
                raise ValueError(f"Destatis incomplete district key at physical row {row_number}")
            district_codes.append(destatis_district_code(row[2], row[3], row[4]))
    if len(district_codes) != 401 or len(set(district_codes)) != 401:
        raise ValueError(
            f"Destatis expected 401 unique district codes, observed rows={len(district_codes)}, unique={len(set(district_codes))}"
        )
    header_cells = {
        f"{chr(64 + column)}{row}": sheet.cell(row, column).value
        for row in range(3, 6)
        for column in range(1, 21)
        if sheet.cell(row, column).value is not None
    }
    return {
        "source": "destatis",
        "reference_date": reference_date,
        "sheet": sheet_name,
        "title_cell": "A1",
        "title": title,
        "physical_header_range": "A3:T5",
        "blank_physical_row": 6,
        "data_first_physical_row": 7,
        "physical_schema": DESTATIS_COLUMNS,
        "observed_nonblank_header_cells": header_cells,
        "row_grain": "one accepted published hierarchy row per source revision",
        "natural_key": ["source_revision_id", "Satzart", "ARS components", "Gemeindename"],
        "required_columns_present": True,
        "accepted_hierarchy_rows": accepted_rows,
        "satzart_distribution": dict(sorted(satzart_counts.items())),
        "trailing_metadata_rows": trailing_metadata,
        "district_rows": len(district_codes),
        "distinct_district_codes": len(set(district_codes)),
        "parsing_ambiguity_changes_grain": False,
    }
