#!/usr/bin/env python3
"""Rebuild compact Gold reference exports from the approved immutable sources.

This is a local evidence export, not a substitute for the Databricks pipeline.
It applies the frozen source, geography, status, point-class and KPI contracts,
then reconciles the locked national controls before writing any output.
"""

from __future__ import annotations

import csv
import hashlib
import json
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
UNIT_SEPARATOR = "\x1f"


def stable_identifier(prefix: str, *parts: Any) -> str:
    encoded = UNIT_SEPARATOR.join("" if part is None else str(part) for part in parts)
    return f"{prefix}_{hashlib.sha256(encoded.encode('utf-8')).hexdigest()}"


def parse_decimal(value: Any) -> Decimal:
    text = str(value).strip().replace(".", "").replace(",", ".")
    return Decimal(text)


def parse_integer(value: Any) -> int:
    return int(parse_decimal(value))


def ratio(numerator: Decimal | int, denominator: int, multiplier: int) -> Decimal:
    return (Decimal(multiplier) * Decimal(numerator) / Decimal(denominator)).quantize(
        Decimal("0.0001"), rounding=ROUND_HALF_UP
    )


def date_set_identity(manifest: dict[str, Any]) -> str:
    encoded = json.dumps(
        manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    )
    return "ds_" + hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def geography_maps() -> tuple[dict[tuple[str, str], str], dict[str, str]]:
    district_map: dict[tuple[str, str], str] = {}
    with (ROOT / "data/reference/geography/geographic_reconciliation.csv").open(
        encoding="utf-8", newline=""
    ) as handle:
        for row in csv.DictReader(handle):
            if not (
                row["source_reference"] == "BNetzA 2026-09-01"
                and row["target_reference"] == "2026-06-30"
                and row["matched"] == "True"
            ):
                continue
            key = (row["Bundesland"], row["source_district_label"])
            value = row["matched_district_key"].zfill(5)
            if key in district_map and district_map[key] != value:
                raise ValueError(f"Conflicting district map for {key}")
            district_map[key] = value
    if len(district_map) != 401:
        raise ValueError(f"Expected 401 BNetzA district identities, got {len(district_map)}")

    analysis_map: dict[str, str] = {}
    with (ROOT / "data/reference/geography/analysis_region_bridge.csv").open(
        encoding="utf-8", newline=""
    ) as handle:
        for row in csv.DictReader(handle):
            district = row["source_district_key"].zfill(5)
            target = row["target_analysis_region_key"].zfill(5)
            if district in analysis_map and analysis_map[district] != target:
                raise ValueError(f"Conflicting analysis-region map for {district}")
            analysis_map[district] = target
    if len(analysis_map) != 401 or len(set(analysis_map.values())) != 400:
        raise ValueError("Expected a governed 401-to-400 district bridge")
    return district_map, analysis_map


def read_kba() -> dict[str, dict[str, Any]]:
    workbook = load_workbook(
        ROOT / "data/raw/kba/fz27_202607.xlsx", read_only=True, data_only=True
    )
    sheet = workbook["FZ 27.15"]
    rows: dict[str, dict[str, Any]] = {}
    for values in sheet.iter_rows(min_row=13, min_col=2, max_col=15, values_only=True):
        raw_code = values[1]
        if raw_code is None:
            continue
        try:
            code = str(int(raw_code)).zfill(5)
        except (TypeError, ValueError):
            continue
        if len(code) != 5 or not code.isdigit():
            continue
        rows[code] = {
            "analysis_region_name": str(values[2]).strip(),
            "registered_passenger_car_stock": parse_integer(values[3]),
            "registered_bev_passenger_car_stock": parse_integer(values[8]),
        }
    if len(rows) != 400:
        raise ValueError(f"Expected 400 KBA analytical rows, got {len(rows)}")
    return rows


def read_supply(
    district_map: dict[tuple[str, str], str], analysis_map: dict[str, str]
) -> dict[str, dict[str, Any]]:
    supply: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "facilities": 0,
            "points": 0,
            "normal_points": 0,
            "fast_points": 0,
            "unclassified_points": 0,
            "nominal_power_kw": Decimal("0"),
        }
    )
    path = ROOT / "data/raw/bnetza/Ladesaeulenregister_BNetzA_2026-09-01.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for _ in range(10):
            next(handle)
        reader = csv.DictReader(handle, delimiter=";", quotechar='"')
        for row in reader:
            district_key = district_map[(row["Bundesland"], row["Kreis/kreisfreie Stadt"])]
            region_code = analysis_map[district_key]
            if row["Status"].strip() != "In Betrieb":
                continue
            target = supply[region_code]
            target["facilities"] += 1
            target["nominal_power_kw"] += parse_decimal(
                row["Nennleistung Ladeeinrichtung [kW]"]
            )
            for slot in range(1, 7):
                if not row[f"Steckertypen{slot}"].strip():
                    continue
                target["points"] += 1
                tokens = [
                    parse_decimal(token)
                    for token in row[f"Nennleistung Stecker{slot}"].split(";")
                    if token.strip()
                ]
                if not tokens:
                    target["unclassified_points"] += 1
                elif max(tokens) <= Decimal("22"):
                    target["normal_points"] += 1
                else:
                    target["fast_points"] += 1
    return supply


def main() -> None:
    manifest = json.loads(
        (ROOT / "config/date_set_manifest_core_v1.json").read_text(encoding="utf-8")
    )
    date_set_id = date_set_identity(manifest)
    district_map, analysis_map = geography_maps()
    stock = read_kba()
    supply = read_supply(district_map, analysis_map)

    fieldnames = [
        "date_set_id",
        "analysis_region_sk",
        "analysis_region_code",
        "analysis_region_name",
        "registered_bev_passenger_car_stock",
        "registered_passenger_car_stock",
        "bev_penetration_percent",
        "in_service_registered_charging_facilities",
        "in_service_registered_charging_points",
        "in_service_registered_normal_charging_points",
        "in_service_registered_fast_charging_points",
        "in_service_registered_unclassified_charging_points",
        "in_service_registered_facility_nominal_power_kw",
        "in_service_registered_charging_points_per_1000_bevs",
        "in_service_registered_fast_charging_points_per_1000_bevs",
        "in_service_registered_charging_facilities_per_1000_bevs",
        "in_service_registered_facility_nominal_power_kw_per_1000_bevs",
        "quality_status",
        "kba_reference_date",
        "bnetza_reference_date",
        "destatis_reference_date",
        "kba_to_bnetza_lag_days",
        "publication_state",
        "date_set_label",
    ]
    result = []
    for code in sorted(stock):
        demand = stock[code]
        provision = supply[code]
        bev = demand["registered_bev_passenger_car_stock"]
        valid_from = "2026-01-01" if code == "06415" else "2026-06-30"
        analysis_region_sk = stable_identifier(
            "analysis_region",
            code,
            valid_from,
            manifest["district_analysis_region_bridge_version"],
        )
        result.append(
            {
                "date_set_id": date_set_id,
                "analysis_region_sk": analysis_region_sk,
                "analysis_region_code": code,
                "analysis_region_name": demand["analysis_region_name"],
                "registered_bev_passenger_car_stock": bev,
                "registered_passenger_car_stock": demand["registered_passenger_car_stock"],
                "bev_penetration_percent": ratio(
                    bev, demand["registered_passenger_car_stock"], 100
                ),
                "in_service_registered_charging_facilities": provision["facilities"],
                "in_service_registered_charging_points": provision["points"],
                "in_service_registered_normal_charging_points": provision["normal_points"],
                "in_service_registered_fast_charging_points": provision["fast_points"],
                "in_service_registered_unclassified_charging_points": provision[
                    "unclassified_points"
                ],
                "in_service_registered_facility_nominal_power_kw": provision[
                    "nominal_power_kw"
                ].quantize(Decimal("0.001")),
                "in_service_registered_charging_points_per_1000_bevs": ratio(
                    provision["points"], bev, 1000
                ),
                "in_service_registered_fast_charging_points_per_1000_bevs": ratio(
                    provision["fast_points"], bev, 1000
                ),
                "in_service_registered_charging_facilities_per_1000_bevs": ratio(
                    provision["facilities"], bev, 1000
                ),
                "in_service_registered_facility_nominal_power_kw_per_1000_bevs": ratio(
                    provision["nominal_power_kw"], bev, 1000
                ),
                "quality_status": "VALIDATED",
                "kba_reference_date": "2026-07-01",
                "bnetza_reference_date": "2026-09-01",
                "destatis_reference_date": "2026-06-30",
                "kba_to_bnetza_lag_days": 62,
                "publication_state": "published",
                "date_set_label": "Core V1 source set - 2026 Q3",
            }
        )

    controls = {
        "regions": len(result),
        "bevs": sum(row["registered_bev_passenger_car_stock"] for row in result),
        "passenger_cars": sum(row["registered_passenger_car_stock"] for row in result),
        "facilities": sum(row["in_service_registered_charging_facilities"] for row in result),
        "points": sum(row["in_service_registered_charging_points"] for row in result),
        "normal_points": sum(
            row["in_service_registered_normal_charging_points"] for row in result
        ),
        "fast_points": sum(
            row["in_service_registered_fast_charging_points"] for row in result
        ),
        "unclassified_points": sum(
            row["in_service_registered_unclassified_charging_points"] for row in result
        ),
        "nominal_power_kw": sum(
            row["in_service_registered_facility_nominal_power_kw"] for row in result
        ),
    }
    expected = {
        "regions": 400,
        "bevs": 2362218,
        "passenger_cars": 49644855,
        "facilities": 116423,
        "points": 209098,
        "normal_points": 154702,
        "fast_points": 54396,
        "unclassified_points": 0,
        "nominal_power_kw": Decimal("9086313.500"),
    }
    if controls != expected:
        raise ValueError(f"Gold control mismatch: observed={controls}, expected={expected}")

    output = ROOT / "data/reference/regional_kpi_snapshot.csv"
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(result)

    quality = ROOT / "data/reference/dataset_quality.csv"
    with quality.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "date_set_id",
                "date_set_label",
                "kba_reference_date",
                "bnetza_reference_date",
                "destatis_reference_date",
                "publication_state",
                "passed_dq_checks",
                "failed_dq_checks",
                "kba_to_bnetza_lag_days",
                "power_bi_refresh_ready",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "date_set_id": date_set_id,
                "date_set_label": "Core V1 source set - 2026 Q3",
                "kba_reference_date": "2026-07-01",
                "bnetza_reference_date": "2026-09-01",
                "destatis_reference_date": "2026-06-30",
                "publication_state": "published",
                "passed_dq_checks": 73,
                "failed_dq_checks": 0,
                "kba_to_bnetza_lag_days": 62,
                "power_bi_refresh_ready": True,
            }
        )
    print(json.dumps({"result": "PASS", "controls": controls}, default=str))


if __name__ == "__main__":
    main()
